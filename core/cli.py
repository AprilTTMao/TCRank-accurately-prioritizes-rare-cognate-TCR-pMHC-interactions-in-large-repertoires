#!/usr/bin/env python3
"""Container entry point; no API, web server, or persistent worker is required."""
import argparse
import csv
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import threading
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from TCR_run import mode_name, positive_int, read_input, select_methods

ASSET_FOLDERS = {
    "PISTE": "PISTE", "PanPep": "PanPep-main", "pMTnet": "pMTnet",
    "TEINet": "TEINet-master", "ERGO-II": "ERGO-II-master", "bertrand": "bertrand-main",
    "TEIM": "TEIM-main", "ATM-TCR": "ATM-TCR-main", "ImRex": "ImRex", "TEPCAM": "TEPCAM",
    "TULIP": "TULIP-TCR", "epiTCR": "epiTCR-main", "iTCep": "iTCep-master",
    "NetTCR-2.2": "NetTCR-2.2-main", "MixTCRpred": "MixTCRpred-main",
}


def write_csv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def target_sheet_name(peptide, mhc, used_names):
    label = str(peptide or "") + "_" + str(mhc or "").removeprefix("HLA-")
    label = re.sub(r"[\\/*?:\[\]]", "", label).strip("'_") or "target"
    name = label[:31].rstrip("'")
    number = 2
    while name.casefold() in used_names:
        suffix = "_" + str(number)
        name = label[:31 - len(suffix)] + suffix
        number += 1
    used_names.add(name.casefold())
    return name


def write_ranking_excel(path, rows, fields):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    groups = {}
    for row in rows:
        groups.setdefault((row.get("Peptide", ""), row.get("MHC", "")), []).append(row)
    if any(len(group) > 1048575 for group in groups.values()):
        raise ValueError("A peptide-HLA group exceeds Excel's row limit; ranking.csv contains the full results.")

    workbook = Workbook(write_only=True)
    names = set()
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="404040")
    for (peptide, mhc), group in groups.items():
        sheet = workbook.create_sheet(target_sheet_name(peptide, mhc, names))
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        header = []
        for index, field in enumerate(fields, 1):
            width = max(len(field), max(len(str(row.get(field) or "")) for row in group)) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(60, max(12, width))
            cell = WriteOnlyCell(sheet, value=field)
            cell.font = header_font
            cell.fill = header_fill
            header.append(cell)
        sheet.append(header)
        # Group rows retain exactly the ordering and columns exported to the CSV.
        for row in group:
            sheet.append([row.get(field) for field in fields])
        sheet.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(fields)), len(group) + 1)
    workbook.save(path)


def export_results(result, destination, mode, methods, elapsed):
    rows = [dict(row) for row in result["ranking_rows"]]
    unscored = 0
    for row in rows:
        row["input_row"] = int(row["pair_id"]) + 1
        row["ranking_status"] = "scored"
        if not row["available_method_count"]:
            row.update(final_score=None, pmhc_rank=None, ranking_status="no_prediction")
            unscored += 1
        elif not math.isfinite(float(row["final_score"])):
            raise ValueError("Non-finite ranking score; results cannot be exported as a successful run.")
    rows.sort(key=lambda row: (row["group_id"], row["ranking_status"] != "scored", -float(row["final_score"] or 0), row["input_row"]))
    # Candidates without any prediction are retained but never assigned a rank.
    counts = {}
    for row in rows:
        if row["ranking_status"] == "scored":
            counts[row["group_id"]] = counts.get(row["group_id"], 0) + 1
            row["pmhc_rank"] = counts[row["group_id"]]
    sequence_fields = ["betaCDR3", "alphaCDR3"]
    if mode == "vregion":
        sequence_fields = ["betaCDR1", "betaCDR2", "betaCDR3", "alphaCDR1", "alphaCDR2", "alphaCDR3"]
    rank_fields = ["input_row", "Peptide", "MHC", *sequence_fields, "pmhc_rank", "pmhc_group_size",
                   "final_score", "available_method_count", "selected_method_count", "ranking_status"]
    write_csv(destination / "ranking.csv", rows, rank_fields)
    write_ranking_excel(destination / "ranking.xlsx", rows, rank_fields)
    details = [dict(row, input_row=int(row["pair_id"]) + 1) for row in result["method_level_rows"]]
    detail_fields = ["input_row", "Peptide", "MHC", *sequence_fields, "method", "availability",
                     "raw_score", "oriented_score", "empirical_p", "evidence", "rank_within_pmhc", "missing_reason"]
    write_csv(destination / "predictor_scores.csv", details, detail_fields)
    failures = dict(result.get("failed_methods", {}))
    for method in methods:
        if not any(row["method"] == method and row.get("raw_score") is not None for row in details):
            failures.setdefault(method, "No usable predictions for these candidates.")
    warnings = []
    if mode == "vregion":
        warnings.append("NetTCR-2.2 and MixTCRpred are auxiliary outputs; their individual fitted feature coefficients are zero.")
    if len(methods) < (15 if mode == "vregion" else 13):
        warnings.append("A predictor subset uses the existing fitted ranker without refitting; its performance is not established by the full-panel evaluation.")
    scored = [row for row in rows if row["ranking_status"] == "scored"]
    if scored and len({row["final_score"] for row in scored}) == 1:
        warnings.append("All scored candidates have tied final scores; tie order follows input order within each target.")
    status = "completed" if not failures and not unscored else "partial"
    if not scored:
        status = "failed"
    summary = {
        "status": status, "mode": mode, "input_rows": len(rows), "predictors_selected": methods,
        "predictors_completed": [method for method in methods if method not in failures],
        "predictors_failed_or_unavailable": failures, "unscored_candidates": unscored,
        "elapsed_seconds": round(elapsed, 2), "ranking_scope": "within peptide-HLA target",
        "calibration": "saved pooled training reference per predictor, not newly sampled target-specific decoys",
        "warnings": warnings,
    }
    (destination / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    return summary


def report_progress(job_dir, stop):
    reported = {}
    while not stop.wait(0.5):
        try:
            state = json.loads((job_dir / "status.json").read_text())
            for name, info in state.get("method_status", {}).items():
                status = info.get("status")
                if status != "queued" and reported.get(name) != status:
                    print(name + ": " + str(status), flush=True)
                    reported[name] = status
        except (OSError, ValueError):
            pass


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--mode", type=mode_name, default="cdr3")
    parser.add_argument("--predictors", nargs="+", default=["all"])
    parser.add_argument("--threads", type=positive_int, default=2)
    parser.add_argument("--keep-logs", action="store_true")
    args = parser.parse_args(argv)
    output = args.output.resolve()
    start = time.monotonic()
    prepared_output = False
    try:
        rows = read_input(args.input, args.mode)
        methods = select_methods(args.predictors, args.mode)
        if output.exists() and any(output.iterdir()):
            raise ValueError("Output directory must be empty.")
        output.mkdir(parents=True, exist_ok=True)
        prepared_output = True
        for name in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS", "TF_NUM_INTRAOP_THREADS", "TCRANK_CPU_THREADS"):
            os.environ[name] = str(args.threads)
        os.environ.update(TCRANK_WEB_METHOD_WORKERS="1", TF_NUM_INTEROP_THREADS="1", CUDA_VISIBLE_DEVICES="")
        from job_runner import run_job
        with tempfile.TemporaryDirectory(prefix="tcrank-") as directory:
            work = Path(directory)
            runtime = work / "runtime"
            assets = runtime / "Predict_methods"
            assets.mkdir(parents=True)
            print("Preparing predictor files...", flush=True)
            for method in methods:
                folder = ASSET_FOLDERS[method]
                shutil.copytree(Path("/opt/tcrank/Predict_methods") / folder, assets / folder)
            (runtime / "runners").symlink_to(ROOT / "core/runners", target_is_directory=True)
            config = json.loads((ROOT / "core/tools_config.json").read_text())
            config["base_dir"] = str(runtime)
            config_path = work / "tools_config.json"
            config_path.write_text(json.dumps(config))
            normalized = work / "input.csv"
            write_csv(normalized, rows, ["alpha", "beta", "peptide", "mhc"])
            job = work / "job"
            stop = threading.Event()
            reporter = threading.Thread(target=report_progress, args=(job, stop), daemon=True)
            reporter.start()
            try:
                result = run_job(backend_dir=ROOT / "core", cfg_path=config_path,
                                 runners_dir=ROOT / "core/runners", inputfiles_script=ROOT / "core/Platform_2_inputfiles.py",
                                 job_dir=job, payload={"mode": args.mode, "methods": methods, "csv_path": str(normalized)})
                summary = export_results(result, output, args.mode, methods, time.monotonic() - start)
                if args.keep_logs or summary["status"] != "completed":
                    shutil.copytree(job / "logs", output / "logs", dirs_exist_ok=True)
            except Exception:
                if (job / "logs").exists():
                    shutil.copytree(job / "logs", output / "logs", dirs_exist_ok=True)
                raise
            finally:
                stop.set()
                reporter.join()
        print("{}: {} candidates; {}/{} predictors returned usable predictions.".format(
            summary["status"].upper(), len(rows), len(summary["predictors_completed"]), len(methods)), flush=True)
        for warning in summary["warnings"]:
            print("NOTE: " + warning, flush=True)
        if summary["status"] != "completed":
            print("See summary.json and logs for unavailable or failed predictions.", file=sys.stderr)
            return 3 if summary["status"] == "partial" else 1
        return 0
    except Exception as error:
        print("ERROR: " + str(error), file=sys.stderr)
        if prepared_output and not (output / "summary.json").exists():
            (output / "summary.json").write_text(json.dumps({"status": "failed", "error": str(error)}, indent=2) + "\n")
        return 1


if __name__ == "__main__":
    sys.exit(main())
