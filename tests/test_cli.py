import csv
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from TCR_run import ARCHIVE, image_parts, main, mode_name, mount_arg, read_input, select_methods
sys.path.insert(0, str(ROOT / "core"))
from output_matching import attach_ids

spec = importlib.util.spec_from_file_location("container_cli", ROOT / "core/cli.py")
container_cli = importlib.util.module_from_spec(spec)
spec.loader.exec_module(container_cli)


class InputTests(unittest.TestCase):
    def test_modes(self):
        self.assertEqual(mode_name("V"), "vregion")
        self.assertEqual(mode_name("cdr"), "cdr3")

    def test_panel_counts(self):
        self.assertEqual(len(select_methods(["all"], "cdr3")), 13)
        self.assertEqual(len(select_methods(["all"], "vregion")), 15)

    def test_subset(self):
        self.assertEqual(select_methods(["panpep,TEINet", "PanPep"], "cdr3"), ["PanPep", "TEINet"])

    def test_invalid_predictors(self):
        with self.assertRaises(ValueError):
            select_methods(["unknown"], "cdr3")
        with self.assertRaises(ValueError):
            select_methods(["MixTCRpred"], "cdr3")

    def test_examples(self):
        self.assertEqual(len(read_input(ROOT / "examples/cdr3_smoke.csv", "cdr3")), 3)
        self.assertEqual(len(read_input(ROOT / "examples/vregion_smoke.csv", "vregion")), 2)

    def test_normalization(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.csv"
            path.write_text("beta_cdr3,peptide,HLA\ncassirssyeqyf,gilgf vftl,A0201\n")
            row = read_input(path, "cdr3")[0]
            self.assertEqual(row["alpha"], "")
            self.assertEqual(row["mhc"], "HLA-A*02:01")
            self.assertEqual(row["peptide"], "GILGFVFTL")

    def test_bad_rows_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.csv"
            for content in (
                "beta_cdr3,peptide,HLA\n,GILGFVFTL,A0201\n",
                "beta_cdr3,peptide,HLA\nCASSX,GILGFVFTL,A0201\n",
                "beta_cdr3,peptide,HLA\nCASS,GILGFVFTL\n",
                "beta_cdr3,peptide,HLA\nCASS,GILGFVFTL,BAD\n",
                "beta_cdr3,beta,peptide,HLA\nCASS,CASS,GILGFVFTL,A0201\n",
            ):
                path.write_text(content)
                with self.subTest(content=content), self.assertRaises(ValueError):
                    read_input(path, "cdr3")

    def test_wrong_mode(self):
        with self.assertRaises(ValueError):
            read_input(ROOT / "examples/vregion_smoke.csv", "cdr3")

    def test_mount_escaping(self):
        value = mount_arg(Path("/tmp/a,b with spaces.csv"), "/input/candidates.csv", True)
        fields = next(csv.reader([value]))
        self.assertIn("src=/tmp/a,b with spaces.csv", fields)
        self.assertIn("readonly", fields)


class EntrypointTests(unittest.TestCase):
    def test_container_dispatches_without_docker(self):
        with patch.dict(os.environ, {"TCRANK_IN_CONTAINER": "1"}), \
             patch.dict(sys.modules, {"cli": container_cli}), \
             patch.object(container_cli, "main", return_value=0) as pipeline, \
             patch("TCR_run.docker_path") as docker:
            result = main(["--input", "input.csv", "--output", "results", "--mode", "v",
                           "--predictors", "panpep,TEINet", "--threads", "3", "--keep-logs"])
            self.assertEqual(result, 0)
            pipeline.assert_called_once_with([
                "--input", "input.csv", "--output", "results", "--mode", "vregion",
                "--threads", "3", "--predictors", "PanPep", "TEINet", "--keep-logs"])
            docker.assert_not_called()

    def test_container_rejects_host_options(self):
        with patch.dict(os.environ, {"TCRANK_IN_CONTAINER": "1"}), \
             patch("TCR_run.docker_path") as docker:
            self.assertEqual(main(["--load-image", "image.tar.gz"]), 2)
            docker.assert_not_called()


class ImageTests(unittest.TestCase):
    def make_parts(self, root):
        (root / "images").mkdir()
        records = []
        for index, content in enumerate((b"first chunk", b"second chunk")):
            name = "images/" + ARCHIVE + ".part-{:02d}".format(index)
            (root / name).write_bytes(content)
            records.append(hashlib.sha256(content).hexdigest() + "  " + name)
        (root / "SHA256SUMS.txt").write_text("\n".join(records) + "\n")

    def test_parts_verified_in_order(self):
        with tempfile.TemporaryDirectory() as directory, patch("TCR_run.ROOT", Path(directory)):
            root = Path(directory)
            self.make_parts(root)
            self.assertEqual([path.read_bytes() for path in image_parts()], [b"first chunk", b"second chunk"])

    def test_corruption_and_missing_parts(self):
        with tempfile.TemporaryDirectory() as directory, patch("TCR_run.ROOT", Path(directory)):
            root = Path(directory)
            self.make_parts(root)
            part = root / "images" / (ARCHIVE + ".part-01")
            part.write_bytes(b"corrupt")
            with self.assertRaisesRegex(ValueError, "checksum mismatch"):
                image_parts()
            part.unlink()
            with self.assertRaisesRegex(ValueError, "missing"):
                image_parts()

    def test_lfs_pointer_has_actionable_error(self):
        with tempfile.TemporaryDirectory() as directory, patch("TCR_run.ROOT", Path(directory)):
            root = Path(directory)
            self.make_parts(root)
            (root / "images" / (ARCHIVE + ".part-00")).write_bytes(b"version https://git-lfs.github.com/spec/v1\n")
            with self.assertRaisesRegex(ValueError, "Git LFS pointers"):
                image_parts()

    def test_manifest_rejects_paths_outside_images(self):
        with tempfile.TemporaryDirectory() as directory, patch("TCR_run.ROOT", Path(directory)):
            root = Path(directory)
            (root / "SHA256SUMS.txt").write_text("0" * 64 + "  ../other\n")
            with self.assertRaisesRegex(ValueError, "Unexpected file"):
                image_parts()

class ExportTests(unittest.TestCase):
    def result(self):
        return {
            "ranking_rows": [
                self.ranking_row(0, available_method_count=0, final_score=0.0),
                self.ranking_row(1),
            ],
            "method_level_rows": [{"pair_id": 1, "method": "PanPep", "raw_score": 0.5}],
            "failed_methods": {},
        }

    def ranking_row(self, pair_id, peptide="GILGFVFTL", mhc="HLA-A*02:01", **changes):
        row = {
            "pair_id": pair_id, "group_id": peptide + "|" + mhc,
            "Peptide": peptide, "MHC": mhc,
            "alphaCDR3": "CAVRDSNYQLIW", "betaCDR3": "CASSIRSSYEQYF",
            "final_score": 1.2, "available_method_count": 1,
            "selected_method_count": 1, "pmhc_group_size": 2,
        }
        row.update(changes)
        return row

    def assert_excel_matches_csv(self, destination):
        with (destination / "ranking.csv").open() as handle:
            reader = csv.DictReader(handle)
            fields = reader.fieldnames
            csv_rows = list(reader)
        workbook = load_workbook(destination / "ranking.xlsx")
        try:
            expected_groups = {}
            for row in csv_rows:
                expected_groups.setdefault((row["Peptide"], row["MHC"]), []).append(row)
            self.assertEqual(len(workbook.worksheets), len(expected_groups))
            seen_groups = set()
            for sheet in workbook:
                table = list(sheet.values)
                self.assertEqual(list(table[0]), fields)
                data = [dict(zip(fields, values)) for values in table[1:]]
                key = (data[0]["Peptide"], data[0]["MHC"])
                self.assertNotIn(key, seen_groups)
                seen_groups.add(key)
                expected = expected_groups[key]
                self.assertEqual(len(data), len(expected))
                for actual_row, expected_row in zip(data, expected):
                    for field in fields:
                        actual = actual_row[field]
                        if field == "final_score" and expected_row[field]:
                            self.assertAlmostEqual(actual, float(expected_row[field]))
                        else:
                            self.assertEqual("" if actual is None else str(actual), expected_row[field])
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
            self.assertEqual(seen_groups, set(expected_groups))
            return workbook.sheetnames
        finally:
            workbook.close()

    def test_unscored_not_ranked(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            summary = container_cli.export_results(self.result(), destination, "cdr3", ["PanPep"], 1)
            self.assertEqual(summary["status"], "partial")
            with (destination / "ranking.csv").open() as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["pmhc_rank"], "1")
            self.assertEqual(rows[1]["pmhc_rank"], "")
            self.assertEqual(rows[1]["ranking_status"], "no_prediction")
            self.assertTrue((destination / "predictor_scores.csv").is_file())
            self.assertEqual(json.loads((destination / "summary.json").read_text())["unscored_candidates"], 1)
            self.assertEqual(self.assert_excel_matches_csv(destination), ["GILGFVFTL_A0201"])

    def test_excel_keeps_pairs_with_shared_sequences_and_sorts_each_target(self):
        result = self.result()
        result["ranking_rows"] = [
            self.ranking_row(0, final_score=1.0, pmhc_group_size=4),
            self.ranking_row(1, "GLCTLVAML", pmhc_group_size=1),
            self.ranking_row(2, mhc="HLA-A*24:02", pmhc_group_size=1),
            self.ranking_row(3, alphaCDR3="CAMSAIMNRDDKIIF", final_score=2.0, pmhc_group_size=4),
            self.ranking_row(4, final_score=1.0, pmhc_group_size=4),
            self.ranking_row(5, available_method_count=0, pmhc_group_size=4),
        ]
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            container_cli.export_results(result, destination, "cdr3", ["PanPep"], 1)
            self.assertEqual(self.assert_excel_matches_csv(destination),
                             ["GILGFVFTL_A0201", "GILGFVFTL_A2402", "GLCTLVAML_A0201"])
            with (destination / "ranking.csv").open() as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([r["input_row"] for r in rows[:4]], ["4", "1", "5", "6"])
            self.assertEqual([r["pmhc_rank"] for r in rows[:4]], ["1", "2", "3", ""])
            self.assertCountEqual([r["input_row"] for r in rows], [str(i) for i in range(1, 7)])

    def test_excel_vregion_has_same_sequence_columns_as_csv(self):
        result = self.result()
        for row in result["ranking_rows"]:
            row.update(betaCDR1="SGHAT", betaCDR2="FQNEAQ", alphaCDR1="TSINN", alphaCDR2="IRSNERE")
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            container_cli.export_results(result, destination, "vregion", ["PanPep"], 1)
            self.assert_excel_matches_csv(destination)
            workbook = load_workbook(destination / "ranking.xlsx")
            try:
                header = next(workbook.active.values)
                self.assertEqual(list(header[3:9]),
                                 ["betaCDR1", "betaCDR2", "betaCDR3", "alphaCDR1", "alphaCDR2", "alphaCDR3"])
            finally:
                workbook.close()

    def test_excel_long_sheet_names_do_not_merge_targets(self):
        result = self.result()
        result["ranking_rows"] = [self.ranking_row(0, "A" * 40 + "G"), self.ranking_row(1, "A" * 40 + "L")]
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            container_cli.export_results(result, destination, "cdr3", ["PanPep"], 1)
            names = self.assert_excel_matches_csv(destination)
            self.assertEqual(len(set(names)), 2)
            self.assertTrue(all(len(name) <= 31 for name in names))
            self.assertTrue(names[1].endswith("_2"))

    def test_sheet_names_remove_invalid_characters_and_avoid_case_collisions(self):
        names = set()
        first = container_cli.target_sheet_name("GILGFVFTL", "HLA-A*02:01", names)
        second = container_cli.target_sheet_name("gilgfvftl", "HLA-A*02:01", names)
        other = container_cli.target_sheet_name("A/B?C[D]", "HLA-A*02:01", names)
        self.assertEqual(first, "GILGFVFTL_A0201")
        self.assertEqual(second, "gilgfvftl_A0201_2")
        for name in (first, second, other):
            self.assertNotRegex(name, r"[\\/*?:\[\]]")

    def test_excel_all_unscored_rows_retained(self):
        result = self.result()
        result["method_level_rows"] = []
        for row in result["ranking_rows"]:
            row["available_method_count"] = 0
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            summary = container_cli.export_results(result, destination, "cdr3", ["PanPep"], 1)
            self.assertEqual(summary["status"], "failed")
            self.assertEqual(summary["unscored_candidates"], 2)
            self.assert_excel_matches_csv(destination)

    def test_no_usable_method_flagged(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self.result()
            result["method_level_rows"] = []
            summary = container_cli.export_results(result, Path(directory), "cdr3", ["PanPep"], 1)
            self.assertIn("PanPep", summary["predictors_failed_or_unavailable"])

    def test_nonfinite_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self.result()
            result["ranking_rows"][1]["final_score"] = float("nan")
            with self.assertRaises(ValueError):
                container_cli.export_results(result, Path(directory), "cdr3", ["PanPep"], 1)


class MappingTests(unittest.TestCase):
    columns = {"beta_col": "CDR3", "peptide_col": "Peptide"}

    def test_reordered_and_filtered_outputs(self):
        inputs = [{"CDR3": "CASSA", "Peptide": "PEPA"}, {"CDR3": "CASSB", "Peptide": "PEPB"}]
        outputs = [dict(inputs[1], score=0.9)]
        mapped = attach_ids(outputs, inputs, [10, 11], self.columns, "PanPep")
        self.assertEqual(mapped[0]["pair_id"], 11)

    def test_deduplicated_outputs(self):
        row = {"CDR3": "CASSA", "Peptide": "PEPA"}
        mapped = attach_ids([row], [row, row], [0, 1], self.columns, "PanPep")
        self.assertEqual([item["pair_id"] for item in mapped], [0, 1])

    def test_unmatched_output_fails(self):
        with self.assertRaises(ValueError):
            attach_ids([{"CDR3": "DIFFERENT", "Peptide": "PEPA"}],
                       [{"CDR3": "CASSA", "Peptide": "PEPA"}], [0], self.columns, "PanPep")

    def test_atmtcr_column_aliases(self):
        mapped = attach_ids([{"CDR3": "CASSA", "Peptide": "PEPA"}],
                            [{"TCR": "CASSA", "Epitope": "PEPA"}], [4], self.columns, "ATM-TCR")
        self.assertEqual(mapped[0]["pair_id"], 4)


if __name__ == "__main__":
    unittest.main()
